import io
import json
import os
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import create_user, get_store
from app.main import app
from app.models import AiSuggestion, BankEntry, Card
from app.services import ai_judge, bank_entries, books, exports, ledger, llm
from app.services.bank_statements import parse_statement
from app.services.store import Store, company_for_user
from db_helpers import fresh_engine
from test_ledger_workflow import KWANGJU, WOORI


CLAUDE = llm.Client(provider="anthropic", api_key="test-key", model="claude-haiku-5-5")


def company():
    ws = ledger.default_workspace()
    ws.year = "2023"
    ws.company = "주식회사 더블락"
    ws.cards = [Card(code="98001", name="광주은행카드", number="2462", bank="99001")]
    return ws


def statements(ws):
    if not (WOORI.exists() and KWANGJU.exists()):
        raise unittest.SkipTest("private bank statement samples are not available")
    _, woori = parse_statement(WOORI.name, WOORI.read_bytes(), ws.banks, "99000")
    _, kwangju = parse_statement(KWANGJU.name, KWANGJU.read_bytes(), ws.banks, "99001")
    return woori + kwangju


def find(transactions, description, amount):
    return next(t for t in transactions if t.description == description and (t.withdrawal or t.deposit) == amount)


def amounts(lines):
    return [(line.account, line.debit, line.credit) for line in lines]


class RuleTests(unittest.TestCase):
    def setUp(self):
        self.ws = company()
        self.txs = statements(self.ws)
        self.by_id = {t.id: t for t in self.txs}
        self.entries: dict[str, BankEntry] = {}
        bank_entries.prefill(self.txs, self.entries, self.ws)

    def test_transfer_between_company_accounts(self):
        out = next(t for t in self.txs if t.withdrawal == 652000)
        incoming = find(self.txs, "(주)더블락", 650000)
        entry = self.entries[out.id]
        self.assertEqual((entry.transfer_to, entry.fee, entry.memo), (incoming.id, 2000, "계좌간이체"))
        self.assertEqual(amounts(bank_entries.lines(entry, out, self.ws, self.by_id)), [
            ("10300", 650000, 0), ("83100", 2000, 0), ("10300", 0, 652000)])
        views = {v.tx_id: v for v in bank_entries.views(self.txs, self.entries, [], self.ws)}
        self.assertEqual(views[incoming.id].covered_by, out.id)  # the deposit is not booked twice
        self.assertEqual(views[out.id].missing, [])

    def test_card_purchase_goes_through_payable(self):
        tx = find(self.txs, "2462 스타벅스 코리아", 5900)
        entry = self.entries[tx.id]
        self.assertEqual((entry.card, entry.vendor), ("98001", "스타벅스 코리아"))
        entry.account = "81200"
        # A check card takes the money at once: expense / 보통예금.
        self.assertEqual(amounts(bank_entries.lines(entry, tx, self.ws, self.by_id)), [("81200", 5900, 0), ("10300", 0, 5900)])
        self.ws.cards[0].payable = True
        lines = bank_entries.lines(entry, tx, self.ws, self.by_id)
        self.assertEqual(amounts(lines), [("81200", 5900, 0), ("25300", 0, 5900), ("25300", 5900, 0), ("10300", 0, 5900)])
        self.assertEqual({(l.account, l.vendor_code) for l in lines if l.account == "25300"}, {("25300", "98001")})

    def test_cancellation_takes_the_original_account(self):
        original = next(t for t in self.txs if t.description == "2462 한국철도공사" and t.withdrawal == 40100 and t.posted_at.startswith("2023-06-11"))
        cancel = find(self.txs, "2462 취소 한국철도공사", 40100)
        self.entries[original.id].account = "81200"
        bank_entries.prefill(self.txs, self.entries, self.ws)
        entry = self.entries[cancel.id]
        self.assertEqual(entry.account, "81200")
        self.assertIn("취소", entry.autofilled["account"])
        self.assertEqual(entry.card, "")  # refunds land in the bank directly
        self.assertEqual(amounts(bank_entries.lines(entry, cancel, self.ws, self.by_id)), [("10300", 40100, 0), ("81200", 0, 40100)])

    def test_receipt_paid_row_is_booked_by_the_receipt(self):
        from test_ledger_workflow import receipt

        tx = find(self.txs, "2462 스타벅스 코리아", 5900)
        r = receipt(merchant="스타벅스 코리아", date="2023-05-24", total=5900, payment="paid", payment_transaction_id=tx.id, bank="99001")
        view = next(v for v in bank_entries.views(self.txs, self.entries, [r], self.ws) if v.tx_id == tx.id)
        # Until the receipt is posted the bank row keeps its own entry, so 보통예금 never has a gap.
        self.assertEqual((view.covered_by, len(view.lines) > 0), ("", True))
        self.assertEqual(view.checks, ["금액 일치", "같은 날", "상호 일치", "영수증 검토 중"])
        r.confirmed = True
        view = next(v for v in bank_entries.views(self.txs, self.entries, [r], self.ws) if v.tx_id == tx.id)
        self.assertEqual((view.covered_by, view.lines), ("R-001", []))

    def test_foreign_card_slip_links_to_next_day_withdrawal(self):
        from app.services import autofill
        from test_ledger_workflow import receipt

        r = receipt(merchant="Google Domains", date="2023-05-18", posting_date="2023-05-18", payment_date="", bank="", total=10967)
        autofill.autofill(r, self.ws, [r], self.txs)
        tx = self.by_id[r.payment_transaction_id]
        self.assertEqual((tx.posted_at[:10], tx.description, r.payment, r.card), ("2023-05-19", "2462 Google Domains", "paid", "98001"))
        self.assertIn("해외 카드", r.autofilled["payment"])
        # A check card's foreign charge is one entry on the withdrawal date, with no payable in between.
        self.assertEqual(r.posting_date, "2023-05-19")
        r.account, r.vat, r.purpose = "85300", "included", "도메인"
        self.assertEqual([(l.date, l.account) for l in ledger.journal(r, self.ws)], [("2023-05-19", "85300"), ("2023-05-19", "10300")])
        other = receipt(id="R-002", merchant="다른 가게", date="2023-05-18", payment_date="", bank="", total=10967)
        autofill.autofill(other, self.ws, [r, other], self.txs)
        self.assertEqual(other.payment_transaction_id, "")  # a later withdrawal needs the same merchant

    def test_closing_moves_profit_and_loss_into_deficit(self):
        capital, interest = find(self.txs, "김태윤", 2000000), find(self.txs, "결산이자", 77)
        rent = find(self.txs, "임대료", 200500)
        for tx, account in ((capital, "33100"), (interest, "90100"), (rent, "81900")):
            self.entries[tx.id].account = account
            self.entries[tx.id].confirmed = True
        lines = books.lines([], self.txs, self.entries, self.ws)
        closing = [(l.account, l.debit, l.credit, l.memo) for l in lines if l.voucher == "결산"]
        self.assertEqual(closing, [
            # 200,500 = 임대료 200,000 + 이체 수수료 500
            ("81900", 0, 200000, "손익계정에 대체"), ("83100", 0, 500, "손익계정에 대체"), ("90100", 77, 0, "손익계정에 대체"),
            ("40000", 0, 77, "수익에서 대체"), ("40000", 200500, 0, "비용에서 대체"),
            ("40000", 0, 200423, "당기순손실 결손금에 대체"), ("37800", 200423, 0, "당기순손익 대체"),
            ("37800", 0, 200423, "이월결손금으로 대체"), ("37600", 200423, 0, "차기이월미처리결손금"),
        ])
        self.assertEqual(sum(l.debit for l in lines), sum(l.credit for l in lines))

    def test_ledger_book_layout(self):
        tx = find(self.txs, "김태윤", 2000000)
        self.entries[tx.id].account, self.entries[tx.id].confirmed = "33100", True
        lines = books.lines([], self.txs, self.entries, self.ws)
        book = load_workbook(io.BytesIO(exports.ledger_book(lines, books.accounts(self.ws, lines), self.ws)))
        self.assertEqual(book.sheetnames, ["0_보통예금(10300)", "1_자본금(33100)"])
        rows = list(book["0_보통예금(10300)"].iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "회사명:주식회사 더블락")
        self.assertEqual(rows[0][6], "계정과목 : [103] 보통예금")
        self.assertEqual(rows[2], ("04-26", "김태윤", "99001", "광주은행", 2000000, None, 2000000))
        self.assertEqual([r[1] for r in rows[3:5]], ["  [ 월         계 ]", "  [ 누         계 ]"])


class FakeClient:
    """Stands in for anthropic.Anthropic: echoes one answer per requested item."""

    def __init__(self, answer):
        self.answer, self.requests = answer, []
        self.messages = SimpleNamespace(create=self.create)

    def create(self, **request):
        self.requests.append(request)
        content = request["messages"][0]["content"]
        text = "".join(block["text"] for block in content if block["type"] == "text") if isinstance(content, list) else content
        items = json.loads(text.split("분류할 항목\n", 1)[1])
        results = [self.answer(item) for item in items]
        return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=json.dumps({"results": results}))])


def answer(item):
    account = "33100" if item.get("description") == "김태윤" else "99999" if item.get("description") == "OTP발급" else "81200"
    return {"id": item["id"], "account": account, "tax": 0, "vendor": "", "purpose": "", "reason": "테스트 판단", "confidence": 0.9}


class AiJudgeTests(unittest.TestCase):
    def test_recommend_sends_chart_as_enum_and_drops_unknown_accounts(self):
        ws = company()
        items = [{"id": "a", "description": "김태윤"}, {"id": "b", "description": "OTP발급"}]
        fake = FakeClient(answer)
        with patch("app.services.llm.anthropic.Anthropic", return_value=fake):
            result = ai_judge.recommend(items, ws, [], CLAUDE)
        self.assertEqual(set(result), {"a"})  # 99999 is not in the chart
        self.assertEqual((result["a"].account, result["a"].model), ("33100", "Claude claude-haiku-5-5"))
        request = fake.requests[0]
        self.assertEqual(request["model"], "claude-haiku-5-5")
        schema = request["extra_body"]["output_config"]["format"]["schema"]
        codes = schema["properties"]["results"]["items"]["properties"]["account"]["enum"]
        self.assertIn("33100", codes)
        self.assertNotIn("10300", codes)  # the bank side is never the model's choice

    def test_skipped_items_are_asked_again(self):
        calls = []

        def skip_first(item):
            calls.append(item["id"])
            return answer(item) if calls.count(item["id"]) > 1 or item["id"] != "b" else {**answer(item), "id": "zz"}

        with patch("app.services.llm.anthropic.Anthropic", return_value=FakeClient(skip_first)):
            result = ai_judge.recommend([{"id": "a", "description": "김태윤"}, {"id": "b", "description": "x"}], company(), [], CLAUDE)
        self.assertEqual(set(result), {"a", "b"})
        self.assertEqual(calls, ["a", "b", "b"])

    def test_same_counterparty_gets_one_account(self):
        items = [{"id": str(i), "type": "통장 거래", "direction": "출금", "description": d} for i, d in
                 enumerate(["2462 스타벅스 코리아", "2462 스타벅스 코리아", "2462 스타벅스 코리아", "2462 투썸플레이스"])]
        result = {"0": AiSuggestion(account="81200", confidence=0.7), "1": AiSuggestion(account="81200", confidence=0.6),
                  "2": AiSuggestion(account="81100", confidence=0.9), "3": AiSuggestion(account="81100", confidence=0.8)}
        out = ai_judge._consistent(items, result)
        self.assertEqual([out[i].account for i in "0123"], ["81200", "81200", "81200", "81100"])
        self.assertIn("같은 거래처 3건 중 2건", out["2"].reason)

    def test_suggestion_never_overrides_a_person(self):
        entry = BankEntry(tx_id="T-1", account="81100")  # typed by a person: no autofilled reason
        ai_judge.apply_to_entry(entry, 10000, AiSuggestion(account="81200", reason="r"))
        self.assertEqual(entry.account, "81100")
        entry = BankEntry(tx_id="T-1")
        ai_judge.apply_to_entry(entry, 11000, AiSuggestion(account="81200", tax=5000, reason="r"))
        self.assertEqual((entry.account, entry.tax), ("81200", 0))  # VAT above 1/11 of the amount is rejected
        self.assertTrue(entry.autofilled["account"].startswith("AI 추천"))


@unittest.skipUnless(KWANGJU.exists(), "private bank statement sample is not available")
class ApiTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.engine = fresh_engine(self.tmp.name)
        user = create_user(self.engine, "owner@example.com", "owner", password="password123")
        self.store = Store(self.engine, company_for_user(self.engine, user.id))
        self.store.save_workspace(company())
        app.dependency_overrides[get_store] = lambda: self.store
        self.client = TestClient(app).__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        app.dependency_overrides.clear()
        self.engine.dispose()
        self.tmp.cleanup()

    def test_import_posts_everything_automatically(self):
        with patch.dict(os.environ, {"AI_JUDGE": "on"}), patch("app.services.llm.anthropic.Anthropic", return_value=FakeClient(answer)):
            imported = self.client.post("/bank-statements", files={"file": (KWANGJU.name, KWANGJU.read_bytes())}, data={"bank": "99001"}).json()
        auto = imported["ai"]["auto"]
        views = self.client.get("/entries").json()
        otp = next(v for v in views if v["transaction"]["description"] == "OTP발급")
        transfers = {v["tx_id"] for v in views if v["requires_person"]}
        # Transfers with no receipt wait for a person; the fake model's unknown account leaves OTP발급 too.
        self.assertEqual(auto["entries_left"], {otp["tx_id"]: "계정", **{t: "영수증 없는 이체 · 사람 확인 필수" for t in transfers}})
        self.assertEqual(auto["entries"], 199 - 1 - len(transfers))
        self.assertTrue({"김태윤", "임대료", "주민세", "최수훈"} <= {v["transaction"]["description"] for v in views if v["requires_person"]})
        self.assertFalse(any(v["requires_person"] for v in views if v["card"] or "이자" in v["transaction"]["description"]))
        self.client.post("/auto/process")
        self.client.patch(f"/entries/{otp['tx_id']}", json={"account": "83100"})
        self.client.post("/auto/process")
        self.assertEqual({v["tx_id"] for v in self.client.get("/entries").json() if not v["confirmed"]}, transfers)
        self.client.post("/entries/bulk-confirm", json={"ids": sorted(transfers)})  # the person confirms
        self.assertTrue(all(v["confirmed"] for v in self.client.get("/entries").json()))
        cash = self.client.get("/ledger/10300").json()
        self.assertEqual((len(cash["rows"]), cash["closing"]), (199, 2013852))  # the statement's closing balance

    def test_receipt_is_posted_automatically_and_takes_over_its_bank_row(self):
        from datetime import date
        from decimal import Decimal
        from app.models import ReceiptExtraction

        with patch.dict(os.environ, {"AI_JUDGE": "on"}), patch("app.services.llm.anthropic.Anthropic", return_value=FakeClient(answer)):
            self.client.post("/bank-statements", files={"file": (KWANGJU.name, KWANGJU.read_bytes())}, data={"bank": "99001"})
        slip = ReceiptExtraction(merchant="스타벅스 코리아", purchased_at=date(2023, 5, 24), subtotal=Decimal(5364), tax=Decimal(536),
                                 total=Decimal(5900), payment_method="체크카드", category_hint="meals", confidence=0.9, line_items=[])
        with patch.dict(os.environ, {"AI_JUDGE": "on"}), patch("app.main.extract_receipt", return_value=slip), \
                patch("app.services.llm.anthropic.Anthropic", return_value=FakeClient(answer)):
            created = self.client.post("/receipts", files={"file": ("slip.pdf", b"%PDF-1.7", "application/pdf")}).json()
        receipt = self.client.get(f"/receipts/{created['id']}").json()
        self.assertTrue(receipt["confirmed"], receipt["missing"])
        self.assertEqual((receipt["card"], receipt["account"]), ("98001", "81200"))
        row = next(v for v in self.client.get("/entries").json() if v["tx_id"] == receipt["payment_transaction_id"])
        self.assertEqual(row["covered_by"], created["id"])
        cash = self.client.get("/ledger/10300").json()
        self.assertEqual([r["credit"] for r in cash["rows"] if r["date"] == "2023-05-24"], [5900, 5900])  # one per bank row, no double

    def test_import_judge_confirm_and_export(self):
        with patch.dict(os.environ, {"AI_JUDGE": "on", "AUTO_CONFIRM": "off"}), patch("app.services.llm.anthropic.Anthropic", return_value=FakeClient(answer)):
            imported = self.client.post("/bank-statements", files={"file": (KWANGJU.name, KWANGJU.read_bytes())}, data={"bank": "99001"}).json()
        self.assertEqual(imported["ai"]["error"], "")
        self.assertGreater(imported["ai"]["judged"], 100)

        entries = {v["transaction"]["description"] + str(v["transaction"]["deposit"]): v for v in self.client.get("/entries").json()}
        capital = entries["김태윤2000000"]
        self.assertEqual(capital["account"], "33100")
        self.assertEqual(capital["autofilled"]["account"], "AI 추천 · 테스트 판단")
        self.assertEqual(capital["ai"]["confidence"], 0.9)

        edited = self.client.patch(f"/entries/{capital['tx_id']}", json={"memo": "자본금 납입"}).json()
        self.assertNotIn("memo", edited["autofilled"])
        self.assertIn("account", edited["autofilled"])  # still shown as a recommendation until confirmed

        done = self.client.post("/entries/bulk-confirm", json={"ids": [capital["tx_id"]]}).json()
        self.assertEqual(done["confirmed"], [capital["tx_id"]])
        self.assertEqual(self.client.patch(f"/entries/{capital['tx_id']}", json={"memo": "x"}).status_code, 409)
        equity = self.client.get("/ledger/33100").json()
        self.assertEqual([(r["memo"], r["credit"]) for r in equity["rows"]], [("자본금 납입", 2000000)])
        book = load_workbook(io.BytesIO(self.client.get("/exports/ledger-book.xlsx").content))
        self.assertIn("1_자본금(33100)", book.sheetnames)

    def test_judge_error_is_reported(self):
        self.client.post("/bank-statements", files={"file": (KWANGJU.name, KWANGJU.read_bytes())}, data={"bank": "99001"})
        with patch.dict(os.environ, {"AI_JUDGE": "on"}), patch("app.services.ai_judge.recommend", side_effect=ai_judge.AiJudgeError("키 없음")):
            response = self.client.post("/ai/judge", json={"ids": []})
        self.assertEqual((response.status_code, response.json()["detail"]), (502, "키 없음"))


if __name__ == "__main__":
    unittest.main()


class PolicyTests(unittest.TestCase):
    def test_precedents_from_an_earlier_ledger(self):
        from app.services import precedents
        from test_ledger_workflow import TEST_DATA

        path = TEST_DATA / "2023계정별원장_주식회사+더블락.xls"
        if not path.exists():
            self.skipTest("private ledger sample is not available")
        found = {(p.text, p.account): p.count for p in precedents.read(path.name, path.read_bytes())}
        self.assertGreaterEqual(found[("스타벅스 코리아", "81200")], 10)  # card prefix dropped, company habit kept
        self.assertIn(("PAYPAL *ITCH IO", "85300"), found)
        self.assertFalse(any(a in ("10300", "25300", "40000") for _, a in found))  # bank, card clearing, closing skipped
        with self.assertRaises(precedents.PrecedentError):
            precedents.read("a.csv", b"x")

    def test_sales_invoice_is_booked_and_collected(self):
        from app.models import SalesInvoice

        ws = company()
        txs = statements(ws)
        invoice = SalesInvoice(id="S-001", issued="2023-06-30", customer="최수훈", supply=270000, tax=27000, item="용역")
        entries: dict[str, BankEntry] = {}
        bank_entries.prefill(txs, entries, ws, [invoice])
        deposit = find(txs, "최수훈", 297000)
        self.assertEqual(entries[deposit.id].account, "10800")
        # An invoice entered after the model guessed still wins; a person's choice does not get replaced.
        late: dict[str, BankEntry] = {}
        bank_entries.prefill(txs, late, ws)
        ai_judge.apply_to_entry(late[deposit.id], 297000, AiSuggestion(account="33100", reason="추정"))
        bank_entries.prefill(txs, late, ws, [invoice])
        self.assertEqual(late[deposit.id].account, "10800")
        late[deposit.id].account, late[deposit.id].autofilled = "33100", {}
        bank_entries.prefill(txs, late, ws, [invoice])
        self.assertEqual(late[deposit.id].account, "33100")
        entries[deposit.id].confirmed = True
        lines = books.lines([], txs, entries, ws, [invoice])
        receivable = [(l.date, l.debit, l.credit) for l in lines if l.account == "10800"]
        self.assertEqual(receivable, [("2023-06-30", 297000, 0), ("2023-09-17", 0, 297000)])
        self.assertEqual([(l.account, l.credit) for l in lines if l.voucher == "S-001"][1:], [("41200", 270000), ("25500", 27000)])

    def test_card_slip_vat_follows_company_policy(self):
        from datetime import date
        from decimal import Decimal
        from app.models import ReceiptExtraction
        from app.services import autofill
        from test_ledger_workflow import receipt

        ws = company()
        slip = ReceiptExtraction(merchant="카페", purchased_at=date(2023, 5, 8), subtotal=Decimal(3455), tax=Decimal(345), total=Decimal(3800),
                                 payment_method="체크카드", category_hint="meals", confidence=0.9, line_items=[],
                                 approval_number="12345678", business_number="123-45-67890")
        r = receipt(total=3800, extraction=slip)
        autofill.autofill(r, ws, [r], [])
        self.assertEqual((r.vat, r.tax), ("deductible", 345))  # default: split, as the VAT return claims it
        ai_judge.apply_to_receipt(r, AiSuggestion(account="81100", tax=0, reason="r"), ws.card_vat)
        self.assertEqual(r.vat, "deductible")  # company policy, not the model, decides
        ai_judge.apply_to_receipt(r, AiSuggestion(account="81300", tax=0, reason="r"), ws.card_vat)
        self.assertEqual(r.vat, "included")  # 접대비 is never deductible

        ws.card_vat = "expense"
        r2 = receipt(total=3800, extraction=slip)
        autofill.autofill(r2, ws, [r2], [])
        self.assertEqual((r2.vat, r2.tax), ("included", 0))
        # Switching the policy re-applies it, but never over a person's choice.
        ws.card_vat = "deduct"
        self.assertTrue(autofill.apply_card_vat(r2, ws))
        self.assertEqual((r2.vat, r2.tax), ("deductible", 345))
        typed = receipt(id="R-009", total=3800, extraction=slip, vat="included")
        self.assertFalse(autofill.apply_card_vat(typed, ws))

    def test_recurring_transfers_are_judged_together(self):
        ws = company()
        txs = statements(ws)
        entries: dict[str, BankEntry] = {}
        bank_entries.prefill(txs, entries, ws)
        views = bank_entries.views(txs, entries, [], ws)
        rent = next(v for v in views if v.transaction.description == "박붕수")
        self.assertIn("임대료", " ".join(ai_judge.peers(views)[rent.tx_id]))
        self.assertEqual((rent.fee, rent.requires_person), (500, True))

    def test_old_chart_gets_standard_accounts_and_unsure_answers_wait(self):
        import tempfile
        from app.auth import create_user
        from app.services.store import Store, company_for_user

        tmp = tempfile.TemporaryDirectory()
        engine = fresh_engine(tmp.name)
        store = Store(engine, company_for_user(engine, create_user(engine, "o@example.com", "o", password="password123").id))
        old = company()
        old.accounts = [a for a in old.accounts if a.code in {c for c, *_ in ledger.DEFAULT_ACCOUNTS}]  # a chart saved before 자본금 existed
        store.save_workspace(old)
        self.assertIn("33100", {a.code for a in store.workspace().accounts})

        from app import main
        entry = BankEntry(tx_id="T-1", account="25100", autofilled={"account": "AI 추천 · 자본금 계정이 없어 임시"})
        self.assertTrue(main._unsure(entry.autofilled, AiSuggestion(account="25100", confidence=0.2)))
        self.assertFalse(main._unsure(entry.autofilled, AiSuggestion(account="25100", confidence=0.9)))
        self.assertFalse(main._unsure({"account": "회사 기준"}, AiSuggestion(account="25100", confidence=0.2)))  # a rule or person decided
        engine.dispose()
        tmp.cleanup()



class StatementTests(unittest.TestCase):
    def test_statements_balance_and_agree_with_the_closing(self):
        from app.services import statements as reports

        ws = company()
        txs = statements(ws)
        entries: dict[str, BankEntry] = {}
        bank_entries.prefill(txs, entries, ws)
        for description, amount, account in (("김태윤", 2000000, "33100"), ("결산이자", 77, "90100"), ("임대료", 200500, "81900")):
            e = entries[find(txs, description, amount).id]
            e.account, e.confirmed = account, True
        lines = books.lines([], txs, entries, ws)
        accounts = books.accounts(ws, lines)
        income = reports.income_statement(accounts, lines, ws.year)
        self.assertEqual(income.items[-1].label, "Ⅹ. 당기순손실")
        self.assertEqual(income.items[-1].amount, 77 - 200500)
        sheet = reports.balance_sheet(accounts, lines, ws.year)
        self.assertEqual(sheet.checks, [])
        totals = {i.label: i.amount for i in sheet.items if i.level == 0}
        self.assertEqual(totals["자산"], totals["부채와 자본 총계"])
        self.assertEqual(totals["자산"], 2000000 + 77 - 200500)
        deficit = reports.deficit_statement(accounts, lines, ws.year)
        self.assertEqual((deficit.title[:8], deficit.items[-1].amount), ("결손금처리계산서", 200423))
        trial = reports.trial_balance(accounts, lines)
        self.assertEqual(sum(r["debit_total"] for r in trial), sum(r["credit_total"] for r in trial))
        # Before the closing is posted the balance sheet carries the result itself and still balances.
        open_sheet = reports.balance_sheet(accounts, reports.before_closing(lines), ws.year)
        self.assertEqual(open_sheet.checks, [])

        book = load_workbook(io.BytesIO(exports.closing_report_xlsx(lines, accounts, ws)))
        self.assertEqual(book.sheetnames, ["재무상태표", "손익계산서", "결손금처리계산서", "합계잔액시산표"])
        self.assertEqual(load_workbook(io.BytesIO(exports.income_statement_xlsx(lines, accounts, ws))).sheetnames, ["손익계산서"])

    def test_vat_return_settles_and_the_refund_collects_it(self):
        from app.models import VatReturn

        ws = company()
        txs = statements(ws)
        refund = find(txs, "광산세무서", 17430)
        entries: dict[str, BankEntry] = {}
        bank_entries.prefill(txs, entries, ws)
        ai_judge.apply_to_entry(entries[refund.id], 17430, AiSuggestion(account="13500", reason="환급"))
        vat = VatReturn(id="V-001", period_start="2023-01-01", period_end="2023-06-30", filed="2023-07-25", output_tax=27000, input_tax=44430)
        bank_entries.prefill(txs, entries, ws, vat_returns=[vat])
        self.assertEqual(entries[refund.id].account, "12000")  # the refund collects 미수금
        entries[refund.id].confirmed = True
        # Booked 매입세액 in the period: 30,000 against 44,430 claimed -> 14,430 comes back out of cost.
        from app.models import LedgerLine
        body = [LedgerLine(receipt_id="R-1", voucher="R-1", date="2023-05-08", memo="", vendor="", vendor_code="", account=a, debit=d, credit=c)
                for a, d, c in (("81100", 300000, 0), ("13500", 30000, 0), ("10300", 0, 330000))]
        lines = books.vat_lines(vat, body)
        self.assertEqual([(l.account, l.debit, l.credit) for l in lines], [
            ("83000", 0, 14430), ("13500", 14430, 0), ("25500", 27000, 0), ("13500", 0, 44430), ("12000", 17430, 0)])
        self.assertEqual(sum(l.debit for l in lines), sum(l.credit for l in lines))
        all_lines = books.lines([], txs, entries, ws, vat_returns=[vat])
        receivable = [(l.date, l.debit, l.credit) for l in all_lines if l.account == "12000"]
        self.assertEqual(receivable, [("2023-07-25", 17430, 0), ("2023-11-08", 0, 17430)])



class VatCheckTests(unittest.TestCase):
    def slip(self, rid, merchant, total, tax, business, day="2023-05-24"):
        from datetime import date as d
        from decimal import Decimal
        from app.models import ReceiptExtraction
        from test_ledger_workflow import receipt

        ext = ReceiptExtraction(merchant=merchant, purchased_at=d.fromisoformat(day), subtotal=Decimal(total - tax), tax=Decimal(tax),
                                total=Decimal(total), payment_method="체크카드", category_hint="meals", confidence=0.9, line_items=[],
                                approval_number="1" + rid[-3:], business_number=business)
        return receipt(id=rid, merchant=merchant, date=day, posting_date=day, total=total, extraction=ext, original_tax=tax, account="81100",
                       vat="deductible" if tax else "included", tax=tax, autofilled={"vat": "회사 기준"})

    def test_corporate_zero_vat_is_flagged_but_exempt_sellers_are_not(self):
        from app.services import vat_checks

        self.assertTrue(vat_checks.needs_vat_check(self.slip("R-001", "스타벅스 코리아", 5900, 0, "201-81-21515")))
        self.assertFalse(vat_checks.needs_vat_check(self.slip("R-002", "씨유월계미소점", 5850, 0, "319-19-00134")))  # sole proprietor
        self.assertFalse(vat_checks.needs_vat_check(self.slip("R-003", "한국철도공사", 22300, 0, "314-82-10024")))  # exempt transport
        self.assertFalse(vat_checks.needs_vat_check(self.slip("R-004", "아이파크몰", 85500, 7772, "120-81-85957")))  # VAT shown
        self.assertEqual((vat_checks.estimate(85500), vat_checks.estimate(5900)), (7772, 536))

    def test_schedule_reconciles_and_fixes_unambiguous_slips(self):
        from app.models import ScheduleRow, VatReturn
        from app.services import vat_checks

        receipts = [self.slip("R-001", "스타벅스 코리아", 5900, 0, "201-81-21515"), self.slip("R-002", "고기왕", 6400, 290, "123-45-67890"),
                    self.slip("R-003", "아이파크몰", 85500, 0, "120-81-85957", "2023-06-18")]
        vr = VatReturn(id="V-001", period_start="2023-01-01", period_end="2023-06-30", filed="2023-07-25", input_tax=8598, schedule=[
            ScheduleRow(business_number="201-81-21515", count=1, tax=536), ScheduleRow(business_number="123-45-67890", count=1, tax=290),
            ScheduleRow(business_number="120-81-85957", count=1, tax=7772)])
        result = vat_checks.reconcile(vr, receipts)
        status = {r["business_number"]: r["status"] for r in result["rows"]}
        self.assertEqual(status, {"123-45-67890": "일치", "120-81-85957": "장부 누락", "201-81-21515": "장부 누락"})
        self.assertEqual(vat_checks.fixes_from_schedule(vr, receipts), {"R-001": 536, "R-003": 7772})


class VatCheckApiTests(unittest.TestCase):
    setUp, tearDown = ApiTests.setUp, ApiTests.tearDown

    def test_recheck_and_schedule_upload(self):
        from datetime import date
        from decimal import Decimal
        from app.models import ReceiptExtraction

        slip = ReceiptExtraction(merchant="에이치디씨아이파크몰", purchased_at=date(2023, 6, 18), subtotal=Decimal(85500), tax=Decimal(0),
                                 total=Decimal(85500), payment_method="체크카드", category_hint="other", confidence=0.9, line_items=[],
                                 approval_number="55554444", business_number="120-81-85957")
        with patch("app.main.extract_receipt", return_value=slip):
            created = self.client.post("/receipts", files={"file": ("slip.pdf", b"%PDF-1.7", "application/pdf")}).json()
        self.assertIn("법인 가맹점 영수증인데 부가세 0원 · 다시 읽기 필요", created["notices"])
        with patch.dict(os.environ, {"AI_JUDGE": "on"}), patch("app.services.ai_judge.read_receipt_vat", return_value={"supply": 77728, "tax": 7772, "total": 85500, "note": ""}):
            result = self.client.post("/receipts/vat-check").json()
        self.assertEqual((result["checked"], result["fixed"]), (1, 1))
        r = self.client.get(f"/receipts/{created['id']}").json()
        self.assertEqual((r["vat"], r["tax"], r["original_tax"]), ("deductible", 7772, 7772))
        self.assertEqual(r["notices"], [n for n in r["notices"] if "부가세 0원" not in n])

        vr = self.client.post("/vat-returns", json={"id": "x", "period_start": "2023-01-01", "period_end": "2023-06-30", "filed": "2023-07-25",
                                                    "output_tax": 0, "input_tax": 7772}).json()
        rows = [{"kind": "card", "business_number": "120-81-85957", "name": "아이파크몰", "count": 1, "supply": 77728, "tax": 7772}]
        with patch("app.services.ai_judge.read_vat_schedule", return_value=rows):
            rec = self.client.post(f"/vat-returns/{vr['id']}/schedule", files={"file": ("명세서.pdf", b"%PDF", "application/pdf")}).json()
        self.assertEqual((rec["matched"], rec["filed_total"], rec["rows"][0]["status"]), (1, 7772, "일치"))
        self.assertEqual(self.client.get(f"/vat-returns/{vr['id']}/reconcile").json()["rows"][0]["book_tax"], 7772)

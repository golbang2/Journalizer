"""Offline regression cases for relationship retrieval and the recommendation boundary."""
import io
import json
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from app import main
from app.auth import create_user
from app.models import BankTransaction, EntryView, Precedent, ReceiptExtraction, ReceiptRecord, Vendor
from app.services import ai_judge, ledger, llm, precedents, retrieval
from app.services.store import Store, company_for_user
from db_helpers import fresh_engine


def workspace():
    ws = ledger.default_workspace()
    ws.year = "2026"
    ws.vendors = [Vendor(code="V1", name="한빛상점", aliases=["한빛"], payment_aliases=["김사장"])]
    return ws


def receipt(id="R-1", **changes):
    return ReceiptRecord(**{
        "id": id, "file_name": "receipt.pdf", "uploaded_at": "2026-01-01",
        "merchant": "한빛상점", "vendor_code": "V1", "date": "2026-01-01",
        "total": 20000, "account": "81100", "confirmed": True, "purpose": "직원 간식",
        **changes,
    })


def entry(id="T-1", **changes):
    tx = BankTransaction(id=id, bank="B1", bank_name="은행", posted_at="2026-01-01",
                         description=changes.pop("description", "김사장"),
                         withdrawal=changes.pop("withdrawal", 20000), deposit=changes.pop("deposit", 0))
    return EntryView(**{"tx_id": id, "transaction": tx, "lines": [], "missing": [], "balanced": True,
                        "account": "81100", "confirmed": True, **changes})


def item(**changes):
    return {"id": "new", "type": "영수증", "merchant": "한빛", "direction": "출금",
            "date": "2026-02-01", "amount": 20000, "purpose": "직원 간식", **changes}


def search(receipts=(), entries=(), target=None, ws=None):
    target = target or item()
    return retrieval.retrieve([target], list(entries), list(receipts), ws or workspace())[target["id"]]


def test_alias_graph_finds_relevant_history_beyond_old_sixty_example_limit():
    unrelated = [receipt(f"R-{i + 10}", merchant=f"다른상점{i}", vendor_code="", purpose="", date="2026-01-20") for i in range(70)]
    found = search([*unrelated, receipt()], target=item(merchant="김사장", purpose=""))
    assert [e.source_id for e in found] == ["R-1"]
    assert "등록 거래처·별칭" in found[0].reason


def test_purpose_ranks_above_same_vendor_amount_and_recency_and_keeps_conflicts():
    gift = receipt("gift", purpose="거래처 선물", account="81300", date="2026-01-31")
    snack = receipt("snack", total=25000)
    found = search([gift, snack])
    assert [e.source_id for e in found] == ["snack", "gift"]
    assert {e.account for e in found} == {"81100", "81300"}


def test_unknown_source_no_amount_only_match_and_bounded_context():
    assert not search([receipt(merchant="무관상점", vendor_code="", purpose="")], target=item(purpose=""))
    assert len(search([receipt(f"R-{i}") for i in range(20)])) == retrieval.MAX_EVIDENCE


def test_excludes_unconfirmed_future_self_and_same_underlying_payment():
    candidates = [receipt("draft", confirmed=False), receipt("future", date="2026-03-01"),
                  receipt("new"), receipt("linked", payment_transaction_id="T-own"), receipt("good")]
    found = search(candidates, target=item(payment_transaction_id="T-own"))
    assert [e.source_id for e in found] == ["good"]
    assert not search([receipt(payment_transaction_id="new")])


def test_bank_excludes_opposite_direction_covered_skipped_and_transfers():
    found = search(entries=[entry("deposit", withdrawal=0, deposit=20000), entry("covered", covered_by="R-1"),
                            entry("skip", skip=True), entry("transfer", transfer_to="T-other"), entry("good")])
    assert [e.source_id for e in found] == ["good"]


def test_confirmed_receipt_payment_edge_retrieves_bank_label_without_vendor():
    found = search([receipt(vendor_code="", payment_transaction_id="T-linked")],
                   [entry("T-linked", description="납품대금홍길동", covered_by="R-1")],
                   item(merchant="납품대금홍길동", purpose=""))
    assert [e.source_id for e in found] == ["R-1"]
    assert "연결 은행 거래 T-linked" in found[0].path


def test_ambiguous_alias_is_not_an_identity_edge():
    ws = workspace()
    ws.vendors.append(Vendor(code="V2", name="다른업체", aliases=["한빛"]))
    assert not search([receipt()], target=item(purpose=""), ws=ws)


def test_business_number_conflict_cannot_match_even_with_same_name():
    ext = ReceiptExtraction(merchant="한빛", purchased_at="2026-01-01", subtotal=20000, tax=0, total=20000,
                            payment_method="card", category_hint="other", confidence=1, line_items=[],
                            business_number="123-45-67890", approval_number="1234")
    assert not search([receipt(extraction=ext)], target=item(business_number="999-99-99999"))
    # A repeated approval is one piece of evidence, regardless of how many files were uploaded.
    assert len(search([receipt("a", extraction=ext), receipt("b", extraction=ext)])) == 1


def test_bank_vendor_business_conflict_does_not_fall_back_to_name():
    ws = workspace()
    ws.vendors[0].business = "1234567890"
    assert not search(entries=[entry(vendor_code="V1")],
                      target=item(merchant="김사장", business_number="9999999999", purpose=""), ws=ws)


def test_ledger_preserves_file_sheet_row_and_original_date_and_legacy_still_loads():
    book = Workbook()
    sheet = book.active
    sheet.title = "복리후생비(81100)"
    sheet.append(["날짜", "적요", "코드", "거래처"])
    sheet.append(["01-02", "한빛", "V1", "한빛"])
    sheet.append(["01-03", "한빛", "V1", "한빛"])
    content = io.BytesIO()
    book.save(content)
    past = precedents.read("old.xlsx", content.getvalue())
    assert past[0].count == 2
    assert [(s.file_name, s.sheet, s.row_number, s.date) for s in past[0].sources] == [
        ("old.xlsx", sheet.title, 2, "01-02"), ("old.xlsx", sheet.title, 3, "01-03")]
    ws = workspace()
    ws.precedents = past
    ws.precedent_source = "old.xlsx"
    evidence = search(ws=ws)[0]
    assert evidence.sources == past[0].sources
    assert evidence.date == ""  # The old ledger did not supply a year.
    assert Precedent.model_validate({"text": "한빛", "account": "81100"}).sources == []


def test_sparse_ledger_keeps_real_excel_row_numbers_with_bounded_allocation():
    book = Workbook()
    sheet = book.active
    sheet.title = "복리후생비(81100)"
    for col, value in enumerate(["01-02", "한빛", "V1", "한빛"], start=1):
        sheet.cell(10, col, value)
    content = io.BytesIO()
    book.save(content)
    assert precedents.read("sparse.xlsx", content.getvalue())[0].sources[0].row_number == 10
    with patch("app.services.spreadsheets.MAX_ROWS", 5), pytest.raises(precedents.PrecedentError):
        precedents.read("sparse.xlsx", content.getvalue())


def test_prompt_attaches_only_item_evidence_and_saves_server_selected_sources():
    targets = [item(), item(id="other", merchant="모르는업체", purpose="")]
    history = retrieval.retrieve(targets, [], [receipt()], workspace())
    requests = []

    def answer(client, system, content, schema, **kwargs):
        rows = json.loads(content[0].split("분류할 항목\n", 1)[1])
        requests.extend(rows)
        return {"results": [{"id": row["id"], "account": "81100", "confidence": 0.7,
                              "evidence": [{"source_id": "fabricated"}]} for row in rows]}

    with patch("app.services.llm.complete", side_effect=answer):
        result = ai_judge.recommend(targets, workspace(), history, llm.Client("anthropic", "fake", "fake"))
    assert requests[0]["evidence"][0]["source_id"] == "R-1"
    assert requests[1]["evidence"] == []
    assert result["new"].evidence == history["new"]
    assert result["other"].evidence == []
    assert result["new"].retrieval_version == retrieval.VERSION


def test_company_scoped_ai_flow_persists_evidence_and_reopening_removes_source(tmp_path):
    engine = fresh_engine(str(tmp_path))
    try:
        stores = []
        for index in range(2):
            user = create_user(engine, f"owner{index}@example.com", "owner", password="password123")
            store = Store(engine, company_for_user(engine, user.id))
            store.save_workspace(workspace())
            stores.append(store)
        own, other = stores
        own.save_receipt(receipt("R-own"))
        other.save_receipt(receipt("R-private", purpose="비밀 직원 간식"))
        own.save_receipt(receipt("R-new", confirmed=False, account="", date="2026-02-01"))
        seen = []

        def answer(client, system, content, schema, **kwargs):
            rows = json.loads(content[0].split("분류할 항목\n", 1)[1])
            seen.extend(rows)
            return {"results": [{"id": row["id"], "account": "81100", "tax": 0, "confidence": 0.7} for row in rows]}

        with patch.object(main, "_ensure_profile"), patch.object(main, "_llm", return_value=llm.Client("anthropic", "fake", "fake")), \
             patch.object(main, "ai_enabled", return_value=True), patch.object(main, "auto_confirm", return_value=False), \
             patch("app.services.llm.complete", side_effect=answer):
            main._run_ai(own, receipt_ids={"R-new"})
        saved = own.receipt("R-new").ai
        assert [e.source_id for e in saved.evidence] == ["R-own"]
        assert "R-private" not in json.dumps(seen, ensure_ascii=False)
        assert not own.receipt("R-new").confirmed
        source = own.receipt("R-own")
        source.confirmed = False
        own.save_receipt(source)
        assert not retrieval.retrieve([item()], [], own.receipts(), own.workspace())["new"]
        # Historical recommendation keeps its original context for audit purposes.
        assert own.receipt("R-new").ai.evidence[0].source_id == "R-own"
    finally:
        engine.dispose()

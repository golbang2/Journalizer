import io
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.models import BankTransaction, ReceiptExtraction, ReceiptRecord, Vendor
from app.services import ledger, matching
from app.services.bank_statements import StatementParseError, parse_statement
from app.auth import create_user, get_store
from db_helpers import fresh_engine
from app.services.store import Store, company_for_user

TEST_DATA = Path(__file__).resolve().parents[2] / "test_data"
WOORI = TEST_DATA / "1-1 우리은행 거래내역.xlsx"
KWANGJU = TEST_DATA / "1-2 광주은행 거래내역.xls"


def workspace(year="2023"):
    ws = ledger.default_workspace()
    ws.year = year
    ws.vendors = [Vendor(code="00109", name="홍익경영파트너스", alias="기업김상철")]
    return ws


def receipt(**values):
    base = dict(id="R-001", file_name="a.pdf", uploaded_at="2023-01-01T00:00:00", merchant="가게", date="2023-11-15",
                total=10000, posting_date="2023-11-15", payment_date="2023-11-15", bank="99000")
    return ReceiptRecord(**{**base, **values})


def amounts(lines):
    return [(line.date, line.account, line.debit, line.credit) for line in lines]


class JournalTests(unittest.TestCase):
    def test_same_day_payment_clears_existing_payable(self):
        r = receipt(id="R-031", merchant="크몽", total=4765000, account="85300", vat="included", payment="paid", payable_offset=144790, posting_date="2023-11-07", payment_date="2023-11-07")
        self.assertEqual(amounts(ledger.journal(r, workspace())), [
            ("2023-11-07", "85300", 4620210, 0),
            ("2023-11-07", "25100", 144790, 0),
            ("2023-11-07", "10300", 0, 4765000),
        ])

    def test_later_payment_posts_payable_then_settlement(self):
        r = receipt(id="R-044", total=54140, account="83100", vat="deductible", tax=4922, payment="paid", payment_date="2023-11-19")
        lines = ledger.journal(r, workspace())
        self.assertEqual(amounts(lines), [
            ("2023-11-15", "83100", 49218, 0),
            ("2023-11-15", "13500", 4922, 0),
            ("2023-11-15", "25100", 0, 54140),
            ("2023-11-19", "25100", 54140, 0),
            ("2023-11-19", "10300", 0, 54140),
        ])
        self.assertEqual({line.voucher for line in lines}, {"R-044-01", "R-044-02"})
        self.assertEqual(lines[-1].vendor, "우리은행")

    def test_missing_lists_every_required_review_item(self):
        gaps = ledger.missing(receipt(), workspace())
        for item in ("사용 목적", "거래처 코드", "비용 계정", "부가세 처리 판단", "증빙·부가세 검토 근거", "지급 상태", "원본·거래 내역 확인"):
            self.assertIn(item, gaps)
        complete = receipt(purpose="음료", vendor_code="S001", account="81100", vat="included", evidence="카드 영수증",
                           payment="paid", payment_link="우리은행 출금", checked=True)
        self.assertEqual(ledger.missing(complete, workspace()), [])
        self.assertIn("회계연도에 맞는 전기일", ledger.missing(complete, workspace("2024")))
        self.assertIn("회사 계정표에 있는 비용 계정", ledger.missing(complete.model_copy(update={"account": "10300"}), workspace()))

    def test_unlinked_bank_payment_is_a_notice_not_a_blocker(self):
        r = receipt(purpose="음료", vendor_code="S001", account="81100", vat="included", evidence="영수증", payment="paid", checked=True)
        view = ledger.view(r, workspace())
        self.assertEqual(view.missing, [])
        self.assertTrue(view.notices[0].startswith("출금 거래 미연결"))
        linked = ledger.view(r.model_copy(update={"payment_link": "우리은행 출금"}), workspace())
        self.assertEqual(linked.notices, [])

    def test_ledger_only_includes_confirmed_complete_receipts(self):
        done = receipt(purpose="음료", vendor_code="S001", account="81100", vat="included", evidence="영수증", payment="paid",
                       payment_link="출금", checked=True, confirmed=True, total=3800)
        draft = done.model_copy(update={"id": "R-002", "confirmed": False})
        stale = done.model_copy(update={"id": "R-003", "purpose": ""})
        ws = workspace()
        view = ledger.ledger(next(a for a in ws.accounts if a.code == "81100"), [done, draft, stale], ws)
        self.assertEqual([(row.voucher, row.debit, row.balance) for row in view.rows], [("R-001-01", 3800, 3800)])

    def test_account_suggestion_prefers_vendor_history(self):
        past = receipt(id="R-001", vendor_code="S001", account="81300", confirmed=True, confirmed_at="2023-12-01")
        extraction = ReceiptExtraction(merchant="가게", purchased_at=date(2023, 12, 2), subtotal=Decimal(1), tax=Decimal(0), total=Decimal(1),
                                       payment_method="카드", category_hint="meals", confidence=0.9, line_items=[])
        new = receipt(id="R-002", vendor_code="S001", extraction=extraction)
        self.assertEqual(ledger.suggest_account(new, [past])[0], "81300")
        self.assertEqual(ledger.suggest_account(new, [])[0], "81100")


@unittest.skipUnless(WOORI.exists() and KWANGJU.exists(), "bank statement samples are not available")
class BankStatementTests(unittest.TestCase):
    def test_woori_xlsx_totals_match_statement_footer(self):
        bank, txs = parse_statement(WOORI.name, WOORI.read_bytes(), workspace().banks)
        self.assertEqual(bank.code, "99000")
        self.assertEqual(sum(t.withdrawal for t in txs), 7150508)
        self.assertEqual(sum(t.deposit for t in txs), 7150903)
        payment = next(t for t in txs if t.posted_at.startswith("2023-11-19"))
        self.assertEqual((payment.description, payment.withdrawal, payment.kind), ("기업김상철", 54140, "인터넷"))

    def test_kwangju_xls_reads_merged_columns(self):
        bank, txs = parse_statement(KWANGJU.name, KWANGJU.read_bytes(), workspace().banks)
        self.assertEqual(bank.code, "99001")
        first = txs[0]
        self.assertEqual((first.posted_at, first.description, first.withdrawal, first.balance), ("2023-12-26T17:43:59", "KJB카드결제2462", 300, 2013852))
        self.assertEqual(first.account_number, "1107-021-555004")

    def test_statements_propose_company_settings(self):
        from app.services import workspace_setup
        from app.services.bank_statements import read_statement

        statements = [(p.name, read_statement(p.name, p.read_bytes())) for p in (WOORI, KWANGJU)]
        proposal = workspace_setup.propose(statements, ledger.default_workspace())
        self.assertEqual((proposal["company"], proposal["year"]), ("주식회사 더블락", "2023"))
        banks = {b["name"]: b for b in proposal["banks"]}
        self.assertEqual((banks["우리은행"]["code"], banks["우리은행"]["number"], banks["우리은행"]["existing"]), ("99000", "1005604530117", True))
        self.assertEqual((banks["광주은행"]["number"], banks["광주은행"]["product"], banks["광주은행"]["opening"]), ("1107-021-555004", "보통예금", 0))
        vendors = {v["name"]: v for v in proposal["vendors"]}
        self.assertTrue(vendors["주식회사 크몽"]["selected"])
        self.assertEqual((vendors["스타벅스 코리아"]["account"], vendors["한국철도공사"]["account"]), ("81100", "81200"))
        self.assertEqual(vendors["PAYPAL *ITCH IO"]["count"], 32)  # truncated variants folded together
        self.assertFalse(vendors["주민세"]["selected"])  # transfer memo, not a vendor
        self.assertFalse(vendors["김태윤"]["selected"])  # deposits only
        self.assertNotIn("예금결산이자", vendors)

    def test_unknown_bank_is_rejected(self):
        with self.assertRaises(StatementParseError):
            parse_statement("statement.xlsx", WOORI.read_bytes(), [])


class MatchingTests(unittest.TestCase):
    def test_exact_amount_with_vendor_alias_ranks_first(self):
        r = receipt(id="R-044", merchant="홍익경영파트너스", vendor_code="00109", total=54140)
        txs = [
            BankTransaction(id="a", bank="99000", bank_name="우리은행", posted_at="2023-11-19T21:15:30", description="기업김상철", withdrawal=54140),
            BankTransaction(id="b", bank="99000", bank_name="우리은행", posted_at="2023-11-15T10:00:00", description="다른곳", withdrawal=54140),
            BankTransaction(id="c", bank="99000", bank_name="우리은행", posted_at="2023-11-15T10:00:00", description="기업김상철", withdrawal=54000),
            BankTransaction(id="d", bank="99000", bank_name="우리은행", posted_at="2024-03-01T10:00:00", description="기업김상철", withdrawal=54140),
        ]
        found = matching.candidates(r, txs, [r], workspace())
        self.assertEqual([c.transaction.id for c in found], ["a", "b"])
        self.assertIn("거래처명 일치", found[0].reasons)


class ApiWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.engine = fresh_engine(self.tmp.name)
        user = create_user(self.engine, "owner@example.com", "owner", password="password123")
        self.store = Store(self.engine, company_for_user(self.engine, user.id))
        self.store.save_workspace(workspace())
        app.dependency_overrides[get_store] = lambda: self.store
        # One event loop for the whole test: each new loop opens a socket pair, which fails
        # intermittently on some Windows setups (WinError 10014).
        self.client = TestClient(app).__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        app.dependency_overrides.clear()
        self.engine.dispose()
        self.tmp.cleanup()

    def extraction(self):
        return ReceiptExtraction(merchant="홍익경영파트너스", purchased_at=date(2023, 11, 15), subtotal=Decimal(49218), tax=Decimal(4922),
                                 total=Decimal(54140), payment_method="계좌이체", category_hint="other", confidence=0.93, line_items=[])

    def test_upload_review_link_confirm_and_export(self):
        with patch("app.main.extract_receipt", return_value=self.extraction()):
            created = self.client.post("/receipts", files={"file": ("slip.pdf", b"%PDF-1.7", "application/pdf")})
        self.assertEqual(created.status_code, 201, created.text)
        r = created.json()
        self.assertEqual((r["id"], r["total"], r["vendor_code"], r["tax"]), ("R-001", 54140, "00109", 4922))
        self.assertEqual(self.client.get("/receipts/R-001/file").content, b"%PDF-1.7")

        if WOORI.exists():
            imported = self.client.post("/bank-statements", files={"file": (WOORI.name, WOORI.read_bytes())}).json()
            self.assertEqual((imported["bank"], imported["added"]), ("99000", 47))
            again = self.client.post("/bank-statements", files={"file": (WOORI.name, WOORI.read_bytes())}).json()
            self.assertEqual((again["added"], again["duplicates"]), (0, 47))
            candidates = self.client.get("/receipts/R-001/payment-candidates").json()
            tx_id = candidates[0]["transaction"]["id"]
            linked = self.client.patch("/receipts/R-001", json={"payment_transaction_id": tx_id}).json()
            self.assertEqual((linked["payment"], linked["payment_date"], linked["bank"]), ("paid", "2023-11-19", "99000"))
        else:
            self.client.patch("/receipts/R-001", json={"payment": "paid", "payment_date": "2023-11-19", "payment_link": "수동"})

        blocked = self.client.post("/receipts/R-001/confirm")
        self.assertEqual(blocked.status_code, 409)
        review = {"purpose": "경영 자문 수수료", "account": "83100", "vat": "deductible", "evidence": "세금계산서 확인", "checked": True}
        ready = self.client.patch("/receipts/R-001", json=review).json()
        self.assertEqual(ready["missing"], [])
        self.assertTrue(self.client.post("/receipts/R-001/confirm").json()["confirmed"])
        self.assertEqual(self.client.patch("/receipts/R-001", json={"purpose": "x"}).status_code, 409)
        self.assertEqual(self.client.delete("/receipts/R-001").status_code, 409)

        cash = self.client.get("/ledger/10300").json()
        self.assertEqual([(row["date"], row["credit"]) for row in cash["rows"]], [("2023-11-19", 54140)])
        self.assertEqual(cash["closing"], -54140)
        csv_text = self.client.get("/ledger/83100/export.csv").content.decode("utf-8-sig")
        self.assertIn("경영 자문 수수료", csv_text)
        book = load_workbook(io.BytesIO(self.client.get("/exports/workbook.xlsx").content))
        self.assertIn("분개장", book.sheetnames)
        self.assertEqual(book["분개장"].max_row, 2 + 5)

        self.client.post("/receipts/R-001/reopen")
        self.assertEqual(self.client.get("/ledger/10300").json()["rows"], [])

    def test_ocr_failure_still_creates_receipt_for_manual_entry(self):
        with patch("app.main.extract_receipt", side_effect=RuntimeError("no credentials")):
            created = self.client.post("/receipts", files={"file": ("photo.jpg", b"\xff\xd8", "image/jpeg")})
        body = created.json()
        self.assertEqual(created.status_code, 201)
        self.assertIn("RuntimeError", body["extraction_error"])
        self.assertIn("거래처명", body["missing"])
        self.assertEqual(self.client.post("/receipts", files={"file": ("a.txt", b"x")}).status_code, 415)
        self.client.delete("/receipts/R-001")
        with patch("app.main.extract_receipt", side_effect=RuntimeError("x")):
            self.assertEqual(self.client.post("/receipts", files={"file": ("b.png", b"x")}).json()["id"], "R-002")

    def test_reextract_fills_facts_and_keeps_review_input(self):
        with patch("app.main.extract_receipt", side_effect=RuntimeError("no credentials")):
            self.client.post("/receipts", files={"file": ("slip.pdf", b"%PDF-1.7", "application/pdf")})
        self.client.patch("/receipts/R-001", json={"purpose": "자문료"})
        with patch("app.main.extract_receipt", return_value=self.extraction()) as ocr:
            body = self.client.post("/receipts/R-001/extract").json()
        self.assertEqual(ocr.call_args.kwargs["content"], b"%PDF-1.7")
        self.assertIsNone(body["extraction_error"])
        self.assertEqual((body["merchant"], body["total"], body["date"], body["vendor_code"], body["purpose"]),
                         ("홍익경영파트너스", 54140, "2023-11-15", "00109", "자문료"))

    def card_slip(self, approval, total, day, merchant="(주) 크몽", business="411-86-01799", tax=0, method="체크계좌"):
        return ReceiptExtraction(merchant=merchant, purchased_at=date.fromisoformat(day), subtotal=Decimal(total - tax), tax=Decimal(tax),
                                 total=Decimal(total), payment_method=method, category_hint="other", confidence=0.95,
                                 line_items=[], approval_number=approval, business_number=business)

    def upload(self, extraction, name="slip.pdf"):
        with patch("app.main.extract_receipt", return_value=extraction):
            response = self.client.post("/receipts", files={"file": (name, b"%PDF-1.7", "application/pdf")})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def test_card_slip_autofill_links_bank_and_learns_vendor_defaults(self):
        self.store.add_transactions([
            BankTransaction(id="T-99000-a", bank="99000", bank_name="우리은행", posted_at="2023-10-12T10:00:00", description="크몽", withdrawal=41400),
            BankTransaction(id="T-99000-b", bank="99000", bank_name="우리은행", posted_at="2023-10-18T10:00:00", description="크몽", withdrawal=90000),
        ])
        first = self.upload(self.card_slip("85290016", 41400, "2023-10-12", merchant="주식회사 크몽"))
        self.assertEqual(first["vendor_code"], "A0001")
        self.assertEqual((first["vat"], first["payment"], first["payment_transaction_id"]), ("included", "paid", "T-99000-a"))
        self.assertIn("85290016", first["evidence"])
        self.assertEqual(set(first["missing"]), {"사용 목적", "비용 계정", "원본·거래 내역 확인"})
        vendor = self.client.get("/workspace").json()["vendors"][-1]
        self.assertEqual((vendor["name"], vendor["auto"]), ("주식회사 크몽", True))

        second = self.upload(self.card_slip("83530022", 90000, "2023-10-18", merchant="(주) 크몽"))
        self.assertEqual(second["vendor_code"], "A0001")  # "(주)" and "주식회사" are the same vendor
        credit = self.upload(self.card_slip("39199544", 56920, "2023-09-23", method="일시불"))
        self.assertEqual(credit["payment"], "unpaid")

        self.client.patch(f"/receipts/{first['id']}", json={"purpose": "외주 개발 용역", "account": "85300", "checked": True})
        self.assertEqual(self.client.post(f"/receipts/{first['id']}/confirm").status_code, 200)
        taught = self.client.get(f"/receipts/{second['id']}").json()
        self.assertEqual((taught["purpose"], taught["account"]), ("외주 개발 용역", "85300"))
        self.assertIn("account", taught["autofilled"])

        result = self.client.post("/receipts/bulk-confirm", json={"ids": [second["id"], credit["id"]]}).json()
        self.assertEqual(result["confirmed"], [])  # nobody ticked "원본 확인" yet
        self.client.post("/receipts/bulk-update", json={"ids": [second["id"], credit["id"]], "changes": {"checked": True}})
        result = self.client.post("/receipts/bulk-confirm", json={"ids": [second["id"], credit["id"]]}).json()
        self.assertEqual(result["confirmed"], [second["id"], credit["id"]])

    def test_edit_clears_auto_marker_and_foreign_slip_is_not_deductible(self):
        r = self.upload(self.card_slip("909158", 31037, "2023-10-18", merchant="ITCH.IO - GAME STORE", business=None))
        self.assertEqual(r["vat"], "included")
        self.assertIn("해외", r["autofilled"]["vat"])
        edited = self.client.patch(f"/receipts/{r['id']}", json={"vat": "deductible", "tax": 0}).json()
        self.assertNotIn("vat", edited["autofilled"])

    def test_duplicate_upload_is_flagged_and_removed(self):
        a = self.upload(self.card_slip("39199544", 56920, "2023-09-23"))
        b = self.upload(self.card_slip("39199544", 56920, "2023-09-23"))
        self.assertEqual(b["duplicate_of"], a["id"])
        self.assertTrue(b["missing"][0].startswith("중복 영수증"))
        self.assertEqual(self.client.post("/receipts/delete-duplicates").json()["deleted"], [b["id"]])
        self.assertEqual([x["id"] for x in self.client.get("/receipts").json()], [a["id"]])

    def test_workspace_rejects_duplicate_account_codes(self):
        ws = self.client.get("/workspace").json()
        ws["accounts"].append(dict(ws["accounts"][0]))
        self.assertEqual(self.client.put("/workspace", json=ws).status_code, 422)

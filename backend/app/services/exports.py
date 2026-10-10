"""CSV and Excel exports of confirmed vouchers."""
import csv
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models import Account, LedgerLine, LedgerView, Workspace
from app.services.ledger import ledger_view

LEDGER_HEADER = ["날짜", "전표번호", "적요", "거래처 코드", "거래처", "차변", "대변", "잔액"]
JOURNAL_HEADER = ["날짜", "전표번호", "계정 코드", "계정명", "적요", "거래처 코드", "거래처", "차변", "대변"]


def _safe(value):
    """Stop spreadsheet apps from evaluating text as a formula."""
    if isinstance(value, str) and re.match(r"^[=+\-@\t\r]", value):
        return "'" + value
    return value


def _ledger_rows(view: LedgerView, year: str) -> list[list]:
    rows = [[f"{year}-01-01", "", "기초잔액", "", "", "", "", view.account.opening]]
    rows += [[r.date, r.voucher, r.memo, r.vendor_code, r.vendor, r.debit, r.credit, r.balance] for r in view.rows]
    rows.append(["", "", "합계", "", "", view.debit_total, view.credit_total, view.closing])
    return rows


def ledger_csv(view: LedgerView, workspace: Workspace) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow([workspace.company, workspace.year, view.account.code, view.account.name])
    writer.writerow(LEDGER_HEADER)
    for row in _ledger_rows(view, workspace.year):
        writer.writerow([_safe(v) for v in row])
    return ("﻿" + out.getvalue()).encode("utf-8")


def workbook(lines: list[LedgerLine], accounts: list[Account], workspace: Workspace) -> bytes:
    names = {a.code: a.name for a in accounts}
    book = Workbook()
    sheet = book.active
    sheet.title = "분개장"
    sheet.append([f"{workspace.company} {workspace.year} 분개장"])
    sheet.append(JOURNAL_HEADER)
    for line in lines:
        sheet.append([_safe(v) for v in [line.date, line.voucher, line.account, names.get(line.account, ""), line.memo, line.vendor_code, line.vendor, line.debit, line.credit]])
    for account in accounts:
        view = ledger_view(account, lines)
        if not view.rows and not account.opening:
            continue
        ws = book.create_sheet(f"{account.code} {account.name}"[:31])
        ws.append([f"{workspace.company} {workspace.year} 계정별원장 · {account.code} {account.name}"])
        ws.append(LEDGER_HEADER)
        for row in _ledger_rows(view, workspace.year):
            ws.append([_safe(v) for v in row])
    for ws in book.worksheets:
        ws["A1"].font = Font(bold=True)
        for cell in ws[2]:
            cell.font = Font(bold=True)
        for column, width in zip("ABCDEFGHI", [12, 12, 12, 16, 28, 12, 18, 14, 14]):
            ws.column_dimensions[column].width = width
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


BOOK_HEADER = ["날짜", "적    요    란", "코드", "거래처", "차   변", "대   변", "잔   액"]
MONTH_TOTAL = "  [ 월         계 ]"
RUNNING_TOTAL = "  [ 누         계 ]"
OPENING = "  [ 전 기 이 월 ]"


def _blank(value: int):
    return value if value else ""


def ledger_book(lines: list[LedgerLine], accounts: list[Account], workspace: Workspace) -> bytes:
    """계정별원장 in the usual accounting-program layout: one sheet per account, 월계·누계 after each month."""
    book = Workbook()
    book.remove(book.active)
    used = [a for a in accounts if a.opening or any(x.account == a.code for x in lines)]
    bold = Font(bold=True)
    for index, account in enumerate(used):
        view = ledger_view(account, lines)
        ws = book.create_sheet(f"{index}_{account.name}({account.code})"[:31])
        ws.append([_safe(f"회사명:{workspace.company}"), "", "", "", "", "", f"계정과목 : [{account.code[:3]}] {account.name}"])
        ws.append(BOOK_HEADER)
        if account.opening:
            ws.append(["", OPENING, "", "", _blank(account.opening if account.normal == "debit" else 0),
                       _blank(account.opening if account.normal == "credit" else 0), account.opening])
        total_debit = total_credit = 0
        month_debit = month_credit = 0
        for i, row in enumerate(view.rows):
            ws.append([row.date[5:], _safe(row.memo), _safe(row.vendor_code), _safe(row.vendor), _blank(row.debit), _blank(row.credit), _blank(row.balance)])
            month_debit += row.debit
            month_credit += row.credit
            last_of_month = i + 1 == len(view.rows) or view.rows[i + 1].date[:7] != row.date[:7]
            if last_of_month:
                total_debit += month_debit
                total_credit += month_credit
                ws.append(["", MONTH_TOTAL, "", "", _blank(month_debit), _blank(month_credit), ""])
                ws.append(["", RUNNING_TOTAL, "", "", _blank(total_debit), _blank(total_credit), ""])
                month_debit = month_credit = 0
        for cell in ws[1] + ws[2]:
            cell.font = bold
        for row in ws.iter_rows(min_row=3):
            if row[1].value in (MONTH_TOTAL, RUNNING_TOTAL, OPENING):
                for cell in row:
                    cell.font = bold
            for cell in row[4:7]:
                cell.number_format = "#,##0"
        for column, width in zip("ABCDEFG", [8, 34, 9, 26, 14, 14, 14]):
            ws.column_dimensions[column].width = width
    if not book.worksheets:
        book.create_sheet("원장 없음").append(["확정한 전표가 없습니다."])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _statement_sheet(ws, workspace: Workspace, st) -> None:
    """과목 | 세부 금액 | 합계 금액 — the usual two-column statement layout."""
    from openpyxl.styles import Alignment, Border, Side

    bold = Font(bold=True)
    ws.append([st.title])
    ws.append([f"회사명: {workspace.company}", "", "(단위: 원)"])
    ws.append(["과    목", "금    액", ""])
    for item in st.items:
        if item.level == 1:
            ws.append([f"   {item.label}", item.amount, ""])
        else:
            ws.append([item.label, "", item.amount])
            for cell in ws[ws.max_row]:
                cell.font = bold
    for check in st.checks:
        ws.append([f"※ {check}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws["C2"].alignment = Alignment(horizontal="right")
    line = Side(style="thin", color="999999")
    for cell in ws[3]:
        cell.font = bold
        cell.border = Border(top=line, bottom=line)
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=4):
        for cell in row[1:3]:
            cell.number_format = "#,##0;(#,##0)"
    for column, width in zip("ABC", [36, 18, 18]):
        ws.column_dimensions[column].width = width


def income_statement_xlsx(lines: list[LedgerLine], accounts: list[Account], workspace: Workspace) -> bytes:
    from app.services import statements

    book = Workbook()
    ws = book.active
    ws.title = "손익계산서"
    _statement_sheet(ws, workspace, statements.income_statement(accounts, lines, workspace.year))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


TRIAL_HEADER = ["차변 잔액", "차변 합계", "계정과목", "대변 합계", "대변 잔액"]


def closing_report_xlsx(lines: list[LedgerLine], accounts: list[Account], workspace: Workspace) -> bytes:
    """결산보고서: 재무상태표, 손익계산서, 결손금처리계산서(이익잉여금처분계산서), 합계잔액시산표."""
    from app.services import statements

    book = Workbook()
    sheets = [
        ("재무상태표", statements.balance_sheet(accounts, lines, workspace.year)),
        ("손익계산서", statements.income_statement(accounts, lines, workspace.year)),
    ]
    deficit = statements.deficit_statement(accounts, lines, workspace.year)
    sheets.append((deficit.title.split(" ")[0], deficit))
    for index, (title, st) in enumerate(sheets):
        ws = book.active if index == 0 else book.create_sheet()
        ws.title = title
        _statement_sheet(ws, workspace, st)

    ws = book.create_sheet("합계잔액시산표")
    ws.append([f"합계잔액시산표 ({workspace.year}년 12월 31일 현재 · 결산 대체 전)"])
    ws.append([f"회사명: {workspace.company}", "", "", "", "(단위: 원)"])
    ws.append(TRIAL_HEADER)
    rows = statements.trial_balance(accounts, lines)
    for r in rows:
        ws.append([r["debit_balance"] or "", r["debit_total"] or "", _safe(f"{r['name']} ({r['code']})"), r["credit_total"] or "", r["credit_balance"] or ""])
    ws.append([sum(r["debit_balance"] for r in rows), sum(r["debit_total"] for r in rows), "합    계",
               sum(r["credit_total"] for r in rows), sum(r["credit_balance"] for r in rows)])
    ws["A1"].font = Font(bold=True, size=13)
    for cell in ws[3] + ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4):
        for cell in (row[0], row[1], row[3], row[4]):
            cell.number_format = "#,##0"
    for column, width in zip("ABCDE", [15, 15, 26, 15, 15]):
        ws.column_dimensions[column].width = width
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
